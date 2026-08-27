"""Dolu alanlar icin kanit spani ADAYI bulur.

--------------------------------------------------------------------------
NE YAPAR, NE YAPMAZ
--------------------------------------------------------------------------
YAPAR: altin sette ZATEN GIRILMIS bir degeri alir, kaynak metinde o degeri
iceren satiri bulur ve kanit spani olarak onerir.

YAPMAZ: hicbir DEGERE karar vermez. Deger etiketleyicinin kararidir; bu
betik yalnizca "o karar metnin neresinde goruluyor" sorusunu cevaplar.
`regex_extractor` (olculen motor) hicbir yerde cagrilmaz.

--------------------------------------------------------------------------
NEDEN GEREKLI
--------------------------------------------------------------------------
Denetimde olculdu: 100 kaydin yalnizca 45'inde kanit spani var, 270 dolu
alan kanitsiz. Bunun bedeli zaten odendi - 14 kaydin kaynak sayfasi
kampanya rotasyonuyla kayboldu ve o degerlerin NEDEN oyle girildigi artik
yalnizca ekran goruntusunde.

--------------------------------------------------------------------------
BELIRSIZ OLANI ATLAR
--------------------------------------------------------------------------
Deger metinde birden fazla yerde geciyorsa ve hangisinin dogru baglam
oldugu ayirt edilemiyorsa span ONERILMEZ - insana birakilir. Yanlis bir
kanit cumlesi, kanitsiz kalmaktan kotudur: dogrulanmis gorunur ama
dogrulamaz.

Kullanim:
    python gold_dataset/kanit_spani_oner.py            # yalnizca rapor
    python gold_dataset/kanit_spani_oner.py --yaz      # Excel'e yaz
"""

from __future__ import annotations

import ortam_yukle  # noqa: F401 - side effect: .env process ortamina yuklenir

import argparse
import json
import re
import sys
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))

GOLD = KOK / "gold_dataset" / "altin_veri_seti.json"
EXCEL = KOK / "gold_dataset" / "altin_veri_seti.xlsx"

# Alan -> o alandan soz eden anahtar kelimeler. Ayni sayi metinde birden
# fazla yerde gectiginde dogru satiri secmek icin kullanilir.
BAGLAM = {
    "taksit_sayisi": ("taksit",),
    "vade_ay": ("vade",),
    "erteleme_suresi_ay": ("öteleme", "ertelem", "ödemesiz"),
    "finansman_tutari": ("finansman", "limit", "kadar", "arası", "tutar"),
    "odul_miktari": ("kazan", "hediye", "iade", "puan", "lira", "para", "indirim"),
    "kar_payi_orani": ("kâr pay", "kar pay", "kâr oran", "kar oran", "oran"),
}
AYLAR = ("Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz",
         "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık")


def _sayi_bicimleri(deger) -> list[str]:
    """Bir sayinin metinde gorunebilecek yazimlari."""
    if isinstance(deger, float) and deger.is_integer():
        deger = int(deger)
    bicimler = {str(deger)}
    if isinstance(deger, int):
        bicimler.add(f"{deger:,}".replace(",", "."))   # 50000 -> 50.000
        if deger >= 1000 and deger % 1000 == 0:
            bicimler.add(f"{deger // 1000} bin")
    if isinstance(deger, float):
        bicimler.add(str(deger).replace(".", ","))     # 2.99 -> 2,99
    return sorted(bicimler, key=len, reverse=True)


def _tarih_bicimleri(iso: str) -> list[str]:
    """'2026-08-31' -> ['31.08.2026', '31-08-2026', '31 Ağustos 2026', ...]

    KARMA BASAMAK: banka sayfalari gun ve ayi ayni sekilde yazmiyor -
    KT-004'un kaynaginda "13.08.2024 - 1.01.2027" geciyor (gun tek, ay cift
    basamakli). Yalnizca "01.01.2027" ve "1.1.2027" uretilseydi bu tarih
    "metinde bulunamadi" sayilir, etiket dogru oldugu halde kanitsiz
    kalirdi. Basamak ve ayrac kombinasyonlari bu yuzden acik acik uretilir."""
    try:
        y, a, g = (int(x) for x in iso.split("-"))
    except (ValueError, AttributeError):
        return []

    gun_yazimlari = {f"{g:02d}", str(g)}
    ay_yazimlari = {f"{a:02d}", str(a)}
    bicimler = {
        f"{gun}{ayrac}{ay}{ayrac}{y}"
        for gun in gun_yazimlari
        for ay in ay_yazimlari
        for ayrac in (".", "-", "/")
    }
    bicimler.update(f"{gun} {AYLAR[a-1]} {y}" for gun in gun_yazimlari)
    bicimler.update(f"{gun} {AYLAR[a-1]}" for gun in gun_yazimlari)
    return sorted(bicimler, key=len, reverse=True)


def _adaylar(metin: str, bicimler: list[str]) -> list[str]:
    """Bicimlerden birini iceren SATIRLAR."""
    bulunan = []
    for satir in metin.split("\n"):
        sade = " ".join(satir.split())
        if not (12 < len(sade) <= 200):
            continue
        if any(b in sade for b in bicimler):
            bulunan.append(sade)
    return bulunan


def _sec(adaylar: list[str], alan: str) -> tuple[str | None, str]:
    """Tek aday varsa onu; coksa baglam kelimesi iceren TEK adayi secer."""
    if not adaylar:
        return None, "deger metinde bulunamadi"
    if len(adaylar) == 1:
        return adaylar[0], "tek aday"
    anahtarlar = BAGLAM.get(alan, ())
    baglamlilar = [a for a in adaylar
                   if any(k.lower() in a.lower() for k in anahtarlar)]
    if len(baglamlilar) == 1:
        return baglamlilar[0], "baglam kelimesiyle ayirt edildi"
    return None, f"{len(adaylar)} aday, ayirt edilemedi - insana birakildi"


def onerileri_uret() -> list[dict]:
    # TEK KAYNAK: kaydin ham metnini cozen mantik burada KOPYALANMAZ.
    # Olculdu (TEK-025): _ham_kampanyalar() EN GUNCEL snapshot'i verir,
    # tests/test_altin_veri_butunlugu.py ise scraper_kaydini_bul ile
    # BASKA bir snapshot'a bakar. Iki taraf ayrilinca bu betik span'i
    # kabul ediyor, butunluk testi ayni span'i reddediyordu.
    from gold_dataset.excel_to_json import SPAN_VERILEBILIR_ALANLAR, span_metinde_var
    from scraper.scripts.gold_eslesme import scraper_kaydini_bul

    with open(GOLD, encoding="utf-8") as f:
        kayitlar = json.load(f)

    oneriler: list[dict] = []
    for k in kayitlar:
        if k["kayit_id"].startswith(("A-", "B-", "C-", "D-")):
            continue
        try:
            kaynak = scraper_kaydini_bul(k)
        except Exception:  # noqa: BLE001 - eslesme yoksa span onerilemez
            kaynak = None
        if not kaynak:
            continue
        metin = kaynak.get("normalize_metin") or kaynak.get("ham_metin") or ""
        mevcut = k.get("kanit_spanlari") or {}

        for alan in SPAN_VERILEBILIR_ALANLAR:
            deger = k.get(alan)
            if deger in (None, "", []) or alan in mevcut:
                continue
            if alan in ("kampanya_baslangic", "kampanya_bitis"):
                bicimler = _tarih_bicimleri(str(deger))
            elif isinstance(deger, (int, float)):
                bicimler = _sayi_bicimleri(deger)
            else:
                continue  # metin alanlari (hedef_kitle, odul_birimi) elle yazilir
            if not bicimler:
                continue
            cumle, gerekce = _sec(_adaylar(metin, bicimler), alan)
            if cumle and span_metinde_var(cumle, metin):
                oneriler.append({"kayit_id": k["kayit_id"], "alan": alan,
                                 "deger": deger, "span": cumle, "gerekce": gerekce})
            else:
                oneriler.append({"kayit_id": k["kayit_id"], "alan": alan,
                                 "deger": deger, "span": None, "gerekce": gerekce})
    return oneriler


def excele_yaz(oneriler: list[dict]) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(EXCEL)
    sh = wb["2. Altin Veri Seti"]
    basliklar = [c.value for c in sh[1]]
    sutun = basliklar.index("kanit_spanlari") + 1
    satir = {sh.cell(r, 1).value: r for r in range(2, sh.max_row + 1) if sh.cell(r, 1).value}

    yazilan = 0
    for o in oneriler:
        if not o["span"]:
            continue
        r = satir[o["kayit_id"]]
        hucre = sh.cell(r, sutun)
        satirlar = [s for s in (hucre.value or "").split("\n") if s.strip()]
        if any(s.startswith(o["alan"] + ":") for s in satirlar):
            continue  # arada eklenmis olabilir
        satirlar.append(f"{o['alan']}: {o['span']}")
        hucre.value = "\n".join(satirlar)
        hucre.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        yazilan += 1
    wb.save(EXCEL)
    return yazilan


def main() -> None:
    a = argparse.ArgumentParser(description="Dolu alanlar icin kanit spani onerir")
    a.add_argument("--yaz", action="store_true", help="Onerileri Excel'e yaz")
    s = a.parse_args()

    oneriler = onerileri_uret()
    bulunan = [o for o in oneriler if o["span"]]
    atlanan = [o for o in oneriler if not o["span"]]

    print(f"  kanitsiz dolu alan (kaynagi olan kayitlarda): {len(oneriler)}")
    print(f"    span bulundu : {len(bulunan)}")
    print(f"    atlandi      : {len(atlanan)}")

    sebep: dict[str, int] = {}
    for o in atlanan:
        sebep[o["gerekce"].split(" - ")[0].split(",")[0]] = \
            sebep.get(o["gerekce"].split(" - ")[0].split(",")[0], 0) + 1
    if sebep:
        print("\n  Atlanma sebepleri:")
        for g, n in sorted(sebep.items(), key=lambda x: -x[1]):
            print(f"    {g:<40}{n:>4}")

    alan_sayim: dict[str, int] = {}
    for o in bulunan:
        alan_sayim[o["alan"]] = alan_sayim.get(o["alan"], 0) + 1
    print("\n  Bulunan spanlar, alan bazinda:")
    for alan, n in sorted(alan_sayim.items(), key=lambda x: -x[1]):
        print(f"    {alan:<24}{n:>4}")

    if s.yaz:
        n = excele_yaz(oneriler)
        print(f"\n  {n} span Excel'e yazildi.")
        print("  Simdi: python gold_dataset/excel_to_json.py && pytest tests/test_altin_veri_butunlugu.py")
    else:
        print("\n  (Yazmak icin --yaz)")
        for o in bulunan[:8]:
            print(f"    {o['kayit_id']:<8} {o['alan']:<20} {str(o['deger'])[:10]:<11}"
                  f" {o['span'][:60]}")


if __name__ == "__main__":
    main()
