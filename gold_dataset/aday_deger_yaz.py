"""Taslak satirlara ADAY deger yazar - imzalamadan, olcume sokmadan.

--------------------------------------------------------------------------
NE YAPAR, NE YAPMAZ
--------------------------------------------------------------------------
YAPAR: bir okuyucunun kaynak metinden cikardigi aday degerleri, her biri
icin dayandigi CUMLEYLE birlikte Excel'e yazar.

YAPMAZ: `giren_kisi` sutununa dokunmaz. Imzasiz satir olcumun disindadir
(bkz. excel_to_json, "IMZASIZ KAYIT HICBIR IDDIA TASIMAZ"), dolayisiyla
buraya yazilan hicbir deger cikarim motorunun olculdugu referansa
girmez - ta ki bir insan sayfayi gorup imzalayana kadar.

--------------------------------------------------------------------------
SPAN ZORUNLULUGU - bu betigin asil koruma mekanizmasi
--------------------------------------------------------------------------
Her deger icin dayandigi cumle verilmek ZORUNDADIR ve o cumle kaynak
metinde birebir bulunmalidir (span_metinde_var). Bulunamazsa DEGER
YAZILMAZ.

Neden: Kural 4 "deger sayfadan gelir, kafadan degil" der. Span
zorunlulugu bunu mekanik hale getirir - hatirlanan, tahmin edilen ya da
uydurulan bir deger cumleye baglanamayacagi icin dosyaya giremez.

--------------------------------------------------------------------------
BELIRSIZ OLAN YAZILMAZ
--------------------------------------------------------------------------
Metin birden fazla degeri esit olcude destekliyorsa (ornek: DK-010'da
"1.000 TL uzeri 3 taksit; 6.000 TL uzeri 6 taksit; 10.000 TL uzeri 9
taksit") alan BOS birakilir ve gerekce `belirsiz` alanina yazilir.
Yanlis bir deger, bos hucreden kotudur: bos hucre "kaynak sunmuyor" der
ve olculebilir; yanlis deger motoru haksiz cezalandirir.

--------------------------------------------------------------------------
KULLANIM
--------------------------------------------------------------------------
    python gold_dataset/aday_deger_yaz.py adaylar.json          # kontrol
    python gold_dataset/aday_deger_yaz.py adaylar.json --yaz    # Excel'e

Girdi bicimi:
    [
      {"kayit_id": "DK-010",
       "kampanya_turu": "Kart Kampanyasi",
       "alanlar": {"kampanya_bitis": "2026-08-31"},
       "spanlar": {"kampanya_bitis": "Bitiş Tarihi: 31 Ağustos 2026"},
       "belirsiz": {"taksit_sayisi": "metin 3, 6 ve 9 taksiti esit destekliyor"}}
    ]
"""

from __future__ import annotations

import ortam_yukle  # noqa: F401 - side effect: .env process ortamina yuklenir

import argparse
import json
import sys
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
if str(KOK) not in sys.path:
    sys.path.insert(0, str(KOK))

GOLD = KOK / "gold_dataset" / "altin_veri_seti.json"
EXCEL = KOK / "gold_dataset" / "altin_veri_seti.xlsx"

DAMGA = "MAKINE ADAYI - dogrulanmadi, imzasiz, olcum disi"

# Bu betigin asla yazmadigi sutunlar.
YASAK_SUTUNLAR = frozenset({"giren_kisi", "giris_tarihi", "kayit_id",
                            "banka", "kaynak_url", "ekran_goruntusu"})

# Deger yazilabilecek alanlar - SPAN_ISTEMEYEN disindakiler span ister.
#
# maliyet_orani BURADA YOK: excel_to_json'un SPAN_VERILEBILIR_ALANLAR
# listesinde yer almadigi icin ona kanit spani yazilamiyor. Span
# dogrulanamayan bir alana deger yazmak, bu betigin tek koruma
# mekanizmasini o alan icin devre disi birakirdi. Alan gerekirse once
# excel_to_json tarafinda span verilebilir hale getirilmeli.
IZINLI_ALANLAR = frozenset({
    "kar_payi_orani", "oran_periyodu", "vade_ay",
    "finansman_tutari", "odul_miktari", "odul_birimi", "kampanya_avantaji",
    "masraf_durumu", "kampanya_baslangic", "kampanya_bitis", "hedef_kitle",
    "taksit_sayisi", "erteleme_suresi_ay", "kampanya_turu",
})

# Serbest metin alanlari: okuyucunun kendi cumlesiyle ozetledigi alanlar,
# kaynakta birebir aranmaz. Sayisal/olculen alanlar bu listede DEGIL.
#
# oran_periyodu de buradadir ama farkli bir sebeple: degeri sayfadan
# kopyalanan bir metin degil, sabit bir siniflandirmadir (aylik/yillik/
# belirsiz). excel_to_json'un SPAN_VERILEBILIR_ALANLAR listesinde yer
# almadigi icin span yazmak "taninmayan alan" uyarisi uretiyordu.
SPAN_ISTEMEYEN = frozenset({
    "kampanya_avantaji", "kampanya_turu", "hedef_kitle", "oran_periyodu",
})


def _kayitlari_al() -> dict[str, dict]:
    with open(GOLD, encoding="utf-8") as f:
        return {k["kayit_id"]: k for k in json.load(f)}


def dogrula(adaylar: list[dict]) -> tuple[list[dict], list[str]]:
    """Yazilabilir adaylari ve reddedilenlerin gerekcelerini dondurur."""
    # TEK KAYNAK: kaydin ham metnini cozen mantik burada KOPYALANMAZ.
    # Olculdu: bu betik once _ham_kampanyalar() ile EN GUNCEL snapshot'a
    # bakiyordu, tests/test_altin_veri_butunlugu.py ise
    # scraper_kaydini_bul ile BASKA bir snapshot'a. Sonuc: betik span'i
    # kabul ediyor, test ayni span'i reddediyordu (TEK-025). Iki taraf
    # ayni cozumleyiciyi kullanmazsa "arac gecti ama test kirildi"
    # durumu kacinilmazdir.
    from gold_dataset.excel_to_json import span_metinde_var
    from scraper.scripts.gold_eslesme import scraper_kaydini_bul

    kayitlar = _kayitlari_al()
    kabul: list[dict] = []
    ret: list[str] = []

    for aday in adaylar:
        kid = aday.get("kayit_id")
        kayit = kayitlar.get(kid)
        if not kayit:
            ret.append(f"{kid}: altin sette boyle bir kayit yok")
            continue
        if (kayit.get("giren_kisi") or "").strip():
            ret.append(f"{kid}: IMZALI kayit - bu betik imzali satira dokunmaz")
            continue

        try:
            eslesen = scraper_kaydini_bul(kayit) or {}
        except Exception:  # noqa: BLE001 - eslesme yoksa span dogrulanamaz
            eslesen = {}
        metin = eslesen.get("normalize_metin") or eslesen.get("ham_metin") or ""
        alanlar = dict(aday.get("alanlar") or {})
        spanlar = dict(aday.get("spanlar") or {})

        temiz: dict[str, object] = {}
        temiz_span: dict[str, str] = {}
        for alan, deger in alanlar.items():
            if alan in YASAK_SUTUNLAR:
                ret.append(f"{kid}.{alan}: bu sutuna yazilmaz")
                continue
            if alan not in IZINLI_ALANLAR:
                ret.append(f"{kid}.{alan}: izinli alan degil")
                continue
            if alan in SPAN_ISTEMEYEN:
                temiz[alan] = deger
                continue

            span = (spanlar.get(alan) or "").strip()
            if not span:
                ret.append(f"{kid}.{alan}: span verilmedi - deger yazilmaz")
                continue
            if not metin:
                ret.append(f"{kid}.{alan}: korpusta ham metin yok - dogrulanamaz")
                continue
            if not span_metinde_var(span, metin):
                ret.append(f"{kid}.{alan}: span kaynakta bulunamadi - deger yazilmaz")
                continue
            temiz[alan] = deger
            temiz_span[alan] = span

        if temiz or aday.get("belirsiz"):
            kabul.append({"kayit_id": kid, "alanlar": temiz, "spanlar": temiz_span,
                          "belirsiz": dict(aday.get("belirsiz") or {})})
    return kabul, ret


def excele_yaz(kabul: list[dict]) -> int:
    import openpyxl

    wb = openpyxl.load_workbook(EXCEL)
    sh = wb["2. Altin Veri Seti"]
    basliklar = [c.value for c in sh[1]]
    sutun = {b: i + 1 for i, b in enumerate(basliklar)}
    satir = {sh.cell(r, 1).value: r for r in range(2, sh.max_row + 1)
             if sh.cell(r, 1).value}

    yazilan = 0
    for a in kabul:
        r = satir[a["kayit_id"]]
        # Guvenlik agi: imza sutunu bu betikten sonra da bos kalmali.
        imza_oncesi = sh.cell(r, sutun["giren_kisi"]).value

        for alan, deger in a["alanlar"].items():
            sh.cell(r, sutun[alan]).value = deger
            yazilan += 1

        if a["spanlar"]:
            hucre = sh.cell(r, sutun["kanit_spanlari"])
            satirlar = [s for s in (hucre.value or "").split("\n") if s.strip()]
            for alan, span in a["spanlar"].items():
                satirlar = [s for s in satirlar if not s.startswith(alan + ":")]
                satirlar.append(f"{alan}: {span}")
            hucre.value = "\n".join(satirlar)
            hucre.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

        notlar = sh.cell(r, sutun["notlar"])
        parcalar = [DAMGA]
        for alan, gerekce in a["belirsiz"].items():
            parcalar.append(f"BELIRSIZ {alan}: {gerekce}")
        notlar.value = "\n".join(parcalar)
        notlar.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

        assert sh.cell(r, sutun["giren_kisi"]).value == imza_oncesi, \
            f"{a['kayit_id']}: imza sutunu degismis"

    wb.save(EXCEL)
    return yazilan


def main() -> None:
    a = argparse.ArgumentParser(description="Taslak satirlara aday deger yazar")
    a.add_argument("girdi", help="Aday JSON dosyasi")
    a.add_argument("--yaz", action="store_true", help="Excel'e yaz")
    s = a.parse_args()

    with open(s.girdi, encoding="utf-8") as f:
        adaylar = json.load(f)

    kabul, ret = dogrula(adaylar)

    print(f"Aday kayit      : {len(adaylar)}")
    print(f"Yazilabilir     : {len(kabul)}")
    deger_sayisi = sum(len(k["alanlar"]) for k in kabul)
    belirsiz_sayisi = sum(len(k["belirsiz"]) for k in kabul)
    print(f"Yazilacak deger : {deger_sayisi}")
    print(f"Belirsiz birakilan alan: {belirsiz_sayisi}")

    if ret:
        print(f"\nREDDEDILEN ({len(ret)}):")
        for r in ret:
            print(f"  - {r}")

    if s.yaz:
        n = excele_yaz(kabul)
        print(f"\n{n} deger yazildi. giren_kisi sutununa DOKUNULMADI.")
        print("Simdi: python gold_dataset/excel_to_json.py")
        print("Imzalamadan once: python gold_dataset/dogrulama_sayfasi.py")
    else:
        print("\n(Yazmak icin --yaz)")


if __name__ == "__main__":
    main()
