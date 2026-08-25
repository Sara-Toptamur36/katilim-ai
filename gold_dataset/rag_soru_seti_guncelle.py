"""Elle isaretlenmis esdeger kampanyalari RAG soru setine isler.

--------------------------------------------------------------------------
NE YAPAR
--------------------------------------------------------------------------
`rag_esdeger_kampanyalar.xlsx` dosyasindaki `DOGRU_CEVAP_MI = EVET`
satirlarini okur ve o sorunun `beklenen_sluglar` listesine ekler.
Ayrica soruya `esdeger_kaynak: "elle_dogrulandi"` isareti koyar - hangi
sorularin genisletilmis yer gercegi kullandigi denetlenebilsin diye.

--------------------------------------------------------------------------
NEDEN INSAN ONAYI SART (dairesellik yasagi)
--------------------------------------------------------------------------
Korpustaki `kampanya_turu` MAKINE ciktisidir (olculen F1 ~%78). Onu
dogrudan dogru-cevap listesine yazmak, RAG olcumunun yer gercegini
cikarim motorunun kendi ciktisina baglar - altin veri setinin 5. kurali
tam olarak bunu yasaklar.

Bu betik makine onerisini DEGIL, yalnizca insanin EVET dedigi satirlari
alir. `HAYIR` ve BOS birakilmis satirlar dokunulmadan gecilir: bos, "bu
kayda henuz bakilmadi" demektir ve "hayir" ile ayni sey DEGILDIR.

--------------------------------------------------------------------------
GERI ALINABILIR
--------------------------------------------------------------------------
Yazmadan once soru setinin yedegi alinir (`.yedek` uzantisiyla). Betik
idempotenttir: ayni slug ikinci kez eklenmez.

Kullanim:
    python -m gold_dataset.rag_soru_seti_guncelle          # yalnizca rapor
    python -m gold_dataset.rag_soru_seti_guncelle --yaz    # soru setine isle
"""

from __future__ import annotations

import argparse
import collections
import json
import shutil
import sys
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(KOK))

EXCEL = KOK / "gold_dataset" / "rag_esdeger_kampanyalar.xlsx"
SORU_SETI = KOK / "gold_dataset" / "rag_soru_seti.json"

# Zorunlu sutunlar - SIRA DEGIL, BASLIK esas alinir.
#
# Onceki surumde sutun numaralari sabit yazilmisti; uretici betige yeni
# bir sutun eklenince (sayfa_basligi) eslesme SESSIZCE kaydi ve bu betik
# yanlis hucreyi okumaya basladi. Basliktan okumak, iki betigin sutun
# duzeni degistiginde birbirinden ayrisamamasini garanti eder.
ZORUNLU_SUTUNLAR = ("soru", "slug", "DOGRU_CEVAP_MI")


def _sutun_haritasi(sh) -> dict[str, int]:
    harita = {
        (c.value or "").strip(): i + 1
        for i, c in enumerate(sh[1])
        if c.value
    }
    eksik = [b for b in ZORUNLU_SUTUNLAR if b not in harita]
    if eksik:
        raise SystemExit(
            f"{EXCEL.name} beklenen sutunlari tasimiyor: {eksik}. "
            "Listeyi yeniden uret: python -m gold_dataset.rag_esdeger_kampanya_listesi"
        )
    return harita


def isaretlenenleri_oku() -> tuple[dict[str, set[str]], dict[str, int]]:
    """(soru -> EVET denen sluglar, sayaclar)."""
    import openpyxl

    if not EXCEL.exists():
        raise SystemExit(
            f"{EXCEL.name} yok. Once listeyi uret:\n"
            "    python -m gold_dataset.rag_esdeger_kampanya_listesi"
        )

    sh = openpyxl.load_workbook(EXCEL).active
    sutun = _sutun_haritasi(sh)
    evetler: dict[str, set[str]] = collections.defaultdict(set)
    sayac = {"evet": 0, "hayir": 0, "bos": 0}

    for r in range(3, sh.max_row + 1):  # 1: baslik, 2: aciklama satiri
        soru = sh.cell(r, sutun["soru"]).value
        slug = sh.cell(r, sutun["slug"]).value
        karar = (sh.cell(r, sutun["DOGRU_CEVAP_MI"]).value or "").strip().upper()
        if not (soru and slug):
            continue
        if karar == "EVET":
            evetler[soru].add(slug)
            sayac["evet"] += 1
        elif karar in ("HAYIR", "HAYıR", "HAYİR"):
            sayac["hayir"] += 1
        else:
            # BOS = "bakilmadi". "hayir" ile ayni sey degildir; bu ayrim
            # korunur ki yarim kalmis bir isaretleme, tamamlanmis gibi
            # gorunmesin.
            sayac["bos"] += 1
    return dict(evetler), sayac


def soru_setini_guncelle(evetler: dict[str, set[str]], yaz: bool) -> list[tuple[str, int, int]]:
    sorular = json.loads(SORU_SETI.read_text(encoding="utf-8"))
    degisim: list[tuple[str, int, int]] = []

    for kayit in sorular:
        yeni = evetler.get(kayit.get("soru", ""))
        if not yeni:
            continue
        onceki = list(kayit["beklenen_sluglar"])
        birlesik = sorted(set(onceki) | yeni)
        if birlesik == sorted(onceki):
            continue
        degisim.append((kayit["soru"], len(onceki), len(birlesik)))
        if yaz:
            kayit["beklenen_sluglar"] = birlesik
            # PROVENANS: hangi sorunun yer gercegi elle genisletildi.
            # Denetim bunu gorebilmeli - sessizce buyuyen bir dogru-cevap
            # listesi, olcumu aciklanamaz sekilde yukseltirdi.
            kayit["esdeger_kaynak"] = "elle_dogrulandi"

    if yaz and degisim:
        shutil.copy2(SORU_SETI, SORU_SETI.with_suffix(".json.yedek"))
        SORU_SETI.write_text(
            json.dumps(sorular, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
    return degisim


def main() -> None:
    a = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    a.add_argument("--yaz", action="store_true", help="Soru setine gercekten isle")
    secim = a.parse_args()

    evetler, sayac = isaretlenenleri_oku()
    degisim = soru_setini_guncelle(evetler, yaz=secim.yaz)

    print(f"  Excel'de EVET  : {sayac['evet']}")
    print(f"  Excel'de HAYIR : {sayac['hayir']}")
    print(f"  bos (bakilmadi): {sayac['bos']}")
    print()

    if not degisim:
        print("  Soru setinde degisiklik yok.")
        if sayac["bos"] == sayac["evet"] + sayac["hayir"] + sayac["bos"]:
            print("  (Excel henuz doldurulmamis gorunuyor.)")
        return

    print("  Genisleyecek sorular:")
    for soru, onceki, sonraki in degisim:
        print(f"    {soru[:36]:<38}{onceki:>4} -> {sonraki}")
    print()
    if secim.yaz:
        print(f"  Islendi. Yedek: {SORU_SETI.name}.yedek")
        print("  Simdi: python -m scraper.scripts.rag_degerlendirme")
    else:
        print("  (Islemek icin --yaz)")


if __name__ == "__main__":
    main()
