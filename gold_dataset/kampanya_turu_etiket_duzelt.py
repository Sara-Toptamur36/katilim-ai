"""Altin Veri Seti'ndeki `kampanya_turu` etiketlerini semaya baglar.

SORUN (olculdu 25 Agustos 2026): 302 imzali altin kayitta 15 farkli
`kampanya_turu` degeri var, ama api/schemas.py::KampanyaTuru enum'unda
yalnizca 9 deger tanimli. 35 kayit enum DISI bir etiket tasiyor:

    Ticari Kampanya            18
    Musteri Ol Kampanyasi       9
    Sigorta/BES Kampanyasi      2
    Katilma Hesabi Kampanyasi   2
    Yatirim Kampanyasi          2
    POS Kampanyasi              2

NEDEN ONEMLI: cikarim motoru yalnizca enum degerlerini uretebilir. Enum
disi bir gold etiketi, motorun DOGRU calistigi durumda bile hata sayilir
- yani bu 35 kayit olcumde ULASILAMAZ bir tavan yaratir. Bunlar
etiketleyicinin ozensizligi DEGIL: gercek kampanya turleri, sema onlari
tanimiyor (bkz. gold_dataset/kampanya_turu_denetimi.md).

IKI FARKLI VAKA, IKI FARKLI COZUM - karistirilmamali:

  (a) AD KAYMASI: enum'da ZATEN bulunan bir kavram, farkli yazimla
      girilmis. Cozum burasi - etiket kanonik enum degerine cevrilir.
      Bu bir YARGI degil, yeniden adlandirmadir.

  (b) SEMA ACIGI: enum'da karsiligi OLMAYAN gercek bir tur. Cozumu
      etiketi degistirmek DEGIL, enum'u genisletmektir; aksi halde
      gercek bir ayrimi silip veriyi yoksullastiririz.

BU BETIK YALNIZCA (a) VAKASINI UYGULAR. (b) icin bkz. api/schemas.py.

OLCUM ETIGI NOTU: gold'u duzeltmek olculen basariyi da degistirir, bu
yuzden yalnizca KAYNAK METINDEN savunulabilen degisiklikler yapilir ve
her biri asagida gerekcesiyle yazilidir. "Skor yukseldi" bir gerekce
degildir.

Kullanim:
    python -m gold_dataset.kampanya_turu_etiket_duzelt          # rapor
    python -m gold_dataset.kampanya_turu_etiket_duzelt --yaz    # uygula
"""

from __future__ import annotations

import ortam_yukle  # noqa: F401 - side effect: .env process ortamina yuklenir

import argparse
import json
from pathlib import Path

KOK = Path(__file__).resolve().parent.parent
GOLD = KOK / "gold_dataset" / "altin_veri_seti.json"
EXCEL = KOK / "gold_dataset" / "altin_veri_seti.xlsx"
SAYFA = "2. Altin Veri Seti"

# --- (a) AD KAYMASI HARITASI ------------------------------------------------
#
# "Yatirim Kampanyasi" -> "Yatirim Urunu Kampanyasi"
#
# GEREKCE (iki kayit da elle okundu):
#   TF-012 "Yeni Yatirim Hesabiniza Sifir Komisyon Orani"
#   KT-046 "Hisse Senedi Islemleriniz Mil'lere Donussun"
# Ikisi de yatirim URUNU kampanyasi. Enum'daki "Yatirim Urunu
# Kampanyasi" etiketini tasiyan kayitlarla AYNI icerik turu - ornegin
# VK-007 "Hisse Senedi Islemlerinde %75 Komisyon Indirimi" ve HF-003
# "Gumus Islemlerinde Dar Makas Avantaji" o etiketle girilmis. Yani
# ayni kavram iki farkli adla yazilmis; kanonik olan enum'daki.
#
# KT-046 NOTU: kampanya odulu "Mil" oldugu icin "Alisveris Puani
# Kampanyasi" da akla gelebilir. Alinmadi - kampanyanin KONUSU hisse
# senedi islemi, odul yalnizca tesvik. Ayni mantikla VK-007 de yatirim
# sayilmis; iki kaydi ayri siniflara koymak yeni bir tutarsizlik olurdu.
AD_KAYMASI = {
    "Yatirim Kampanyasi": "Yatirim Urunu Kampanyasi",
}


def enum_degerleri() -> set[str]:
    from api.schemas import KampanyaTuru

    return {e.value for e in KampanyaTuru}


def etkilenen_kayitlar() -> tuple[list[dict], list[dict]]:
    """(cevrilecekler, sema acigi olanlar) - ikisi de gerekceleriyle."""
    gold = json.loads(GOLD.read_text(encoding="utf-8"))
    enum = enum_degerleri()
    cevir: list[dict] = []
    sema_acigi: list[dict] = []

    for kayit in gold:
        etiket = kayit.get("kampanya_turu")
        if not etiket or etiket in enum:
            continue
        if etiket in AD_KAYMASI:
            cevir.append({
                "kayit_id": kayit["kayit_id"],
                "eski": etiket,
                "yeni": AD_KAYMASI[etiket],
                "kampanya_adi": kayit.get("kampanya_adi"),
            })
        else:
            sema_acigi.append({
                "kayit_id": kayit["kayit_id"],
                "etiket": etiket,
                "kampanya_adi": kayit.get("kampanya_adi"),
            })

    return cevir, sema_acigi


def excele_yaz(cevir: list[dict]) -> int:
    """Etkilenen satirlarin kampanya_turu hucresini kanonik degere cevirir."""
    import openpyxl

    wb = openpyxl.load_workbook(EXCEL)
    sh = wb[SAYFA]
    basliklar = [c.value for c in sh[1]]
    sutun = {b: i + 1 for i, b in enumerate(basliklar)}
    satir = {
        sh.cell(r, 1).value: r
        for r in range(2, sh.max_row + 1)
        if sh.cell(r, 1).value
    }

    hedef = {c["kayit_id"]: c["yeni"] for c in cevir}
    yazilan = 0
    for kid, yeni in hedef.items():
        r = satir.get(kid)
        if r is None:
            continue
        # Guvenlik agi (vade_farksiz_duzelt.py / aday_deger_yaz.py ile ayni
        # desen): bu betik imza sutununa ASLA dokunmamali - etiketin
        # sorumlusu degismiyor, yalnizca yazimi kanoniklesiyor.
        imza_oncesi = sh.cell(r, sutun["giren_kisi"]).value
        sh.cell(r, sutun["kampanya_turu"]).value = yeni
        assert sh.cell(r, sutun["giren_kisi"]).value == imza_oncesi, \
            f"{kid}: imza sutunu degismis"
        yazilan += 1

    wb.save(EXCEL)
    return yazilan


def main() -> None:
    a = argparse.ArgumentParser(
        description="kampanya_turu etiketlerini KampanyaTuru enum'una baglar"
    )
    a.add_argument("--yaz", action="store_true", help="Excel'e uygula")
    args = a.parse_args()

    cevir, sema_acigi = etkilenen_kayitlar()

    print(f"(a) AD KAYMASI - cevrilecek : {len(cevir)} kayit")
    for c in cevir:
        print(f"    {c['kayit_id']:10} {c['eski']!r} -> {c['yeni']!r}")
        print(f"    {'':10} {c['kampanya_adi']}")

    print()
    print(f"(b) SEMA ACIGI - DOKUNULMAZ : {len(sema_acigi)} kayit")
    turler: dict[str, int] = {}
    for s in sema_acigi:
        turler[s["etiket"]] = turler.get(s["etiket"], 0) + 1
    for t, c in sorted(turler.items(), key=lambda x: -x[1]):
        print(f"    {t:32} {c}")
    print("    Cozum etiketi degistirmek DEGIL, api/schemas.py::KampanyaTuru")
    print("    enum'unu genisletmektir - bu turler gercek ve ayirt edici.")

    if not args.yaz:
        print()
        print("Rapor modu - Excel'e yazmak icin --yaz ekleyin.")
        return

    yazilan = excele_yaz(cevir)
    print()
    print(f"Excel guncellendi: {yazilan} satir")
    print("Simdi JSON'u yeniden uretin:")
    print("    python gold_dataset/excel_to_json.py")


if __name__ == "__main__":
    main()
